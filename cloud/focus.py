# -*- coding: utf-8 -*-
"""专注力分析引擎 + 三件套生成（系统二 ATTENTION · 规则引擎）。

输入流水线结构化转写 segments（含 role/speaker 与时间戳），全部指标规则计算，
不消耗大模型额度；口径与《专注力数据录入表.xlsx · 字段字典》一致：

- 话轮 = 说话人连续发言段；问句 = 老师话轮命中 ？/吗/呢/怎么/为什么/多少/哪个/是不是…
- 应答 = 老师问句后 30 秒内下一段为学生发言
- 长讲解段 = 老师单段 ≥90 秒或 ≥220 字；沉默间隙 = 相邻话轮间隔 >20 秒
- 老师授课分 = 40%问句引导(2.5/min封顶) + 30%讲解节奏(长讲解惩罚) + 30%主题覆盖
- 课堂活跃分 = 40%话轮密度(3/min封顶) + 30%沉默控制 + 30%分段均衡
- 学生专注度 = 40%参与(占比近50%最优) + 30%应答率 + 30%学生分段均衡
- 师生匹配度 = 40%互动同步(交替密度6/min封顶) + 30%参与均衡 + 30%内容契合

输出三件套（与 0820 会议示例一致）：
1. 专注力报告.html（ECharts 内嵌，单文件离线可开）
2. 专注力报告.xlsx（报告概览/课次指标明细/分段活跃度/指标口径与建议）
3. 专注力数据录入表.xlsx（累计事实表，含学生/老师自动汇总公式，逐课追加）
另输出 focus_data.json（引擎原始数据，供回填与复核）。
"""
from __future__ import annotations

import json
import logging
import re
from pathlib import Path

log = logging.getLogger("focus")

ASSETS = Path(__file__).resolve().parent / "assets"

# ---------------------------------------------------------------- 规则常量

_Q_PATTERN = re.compile(r"[？?]|吗[。，、！？? ]|呢|怎么|为什么|为何|多少|几个|哪种|"
                        r"哪个|哪位|是不是|对不对|好不好|行不行|能不能|谁来|试着")
_TOPIC_PATTERN = re.compile(r"[A-Za-z]\d+(?:[/.\-]\d+)*|\d+(?:[/]\d+)+")
_LONG_SEC, _LONG_CHARS = 90, 220
_SILENCE_GAP = 20.0
_ANSWER_WINDOW = 30.0


def _round(x, nd=1):
    return None if x is None else round(x, nd)


# ---------------------------------------------------------------- 1. 指标计算


def compute_focus(segments: list[dict], meta: dict, ctx: dict | None = None) -> dict:
    """由结构化转写计算专注力指标（L schema，与 HTML/xlsx 模板字段一一对应）。"""
    segs = sorted(segments, key=lambda s: s["start_sec"])
    t_segs = [s for s in segs if s.get("role") == "teacher"]
    s_segs = [s for s in segs if s.get("role") == "student"]
    separated = bool(t_segs and s_segs)

    duration_sec = (ctx or {}).get("duration_sec") or meta.get("duration_sec") or 0
    if not duration_sec and segs:
        duration_sec = segs[-1]["end_sec"] - segs[0]["start_sec"]
    duration_min = round(duration_sec / 60, 1) if duration_sec else 0.0

    total_chars = sum(len(s["text"]) for s in segs)
    s_chars = sum(len(s["text"]) for s in s_segs)
    s_ratio = round(s_chars / total_chars * 100, 1) if total_chars else None

    # 话轮交替
    alternations = sum(1 for i in range(1, len(segs))
                       if segs[i]["speaker_id"] != segs[i - 1]["speaker_id"])
    inter_density = round(alternations / duration_min, 2) if duration_min else 0

    # 问句与应答
    q_idx = [i for i, s in enumerate(segs)
             if s.get("role") == "teacher" and _Q_PATTERN.search(s["text"])]
    n_question = len(q_idx)
    answered = 0
    for i in q_idx:
        if i + 1 < len(segs):
            nxt = segs[i + 1]
            if nxt.get("role") == "student" and \
                    nxt["start_sec"] - segs[i]["end_sec"] <= _ANSWER_WINDOW:
                answered += 1
    response_rate = round(answered / n_question * 100) if n_question else None

    # 长讲解（老师）
    longs = [s for s in t_segs
             if s["end_sec"] - s["start_sec"] >= _LONG_SEC or len(s["text"]) >= _LONG_CHARS]
    longest_sec = max((int(s["end_sec"] - s["start_sec"]) for s in t_segs), default=0)

    # 沉默间隙
    gaps = [segs[i]["start_sec"] - segs[i - 1]["end_sec"]
            for i in range(1, len(segs))]
    sil = [g for g in gaps if g > _SILENCE_GAP]
    silence_min = round(sum(sil) / 60, 1)
    silence_ratio = round(sum(sil) / duration_sec * 100, 1) if duration_sec else 0.0

    # 分段活跃（时间三等分，按字数）
    seg_ratio = _thirds_ratio(segs, duration_sec)
    stu_seg_ratio = _thirds_ratio(s_segs, duration_sec) if separated else None
    balance = _balance(seg_ratio)
    stu_balance = _balance(stu_seg_ratio) if stu_seg_ratio else None

    # 主题覆盖（编号/公式特征词去重）
    all_text = "\n".join(s["text"] for s in segs)
    topic_count = len(set(_TOPIC_PATTERN.findall(all_text)))

    # ---- 评分 ----
    q_density = round(n_question / duration_min, 2) if duration_min else 0
    teach_score = _teach_score(q_density, len(longs), longest_sec, topic_count)
    active_score = _active_score(len(segs), duration_min, silence_ratio, balance)

    if separated:
        participation = max(0.0, 100 - abs((s_ratio or 0) - 50) * 2)
        focus_score = round(0.4 * participation
                            + 0.3 * (response_rate or 0)
                            + 0.3 * (stu_balance or 0))
        sync = min(inter_density / 6, 1) * 100
        match_score = round(0.4 * sync + 0.3 * participation
                            + 0.3 * (response_rate or 0))
        focus_est = focus_note = None
    else:
        participation = focus_score = match_score = None
        focus_est = "60-90"
        focus_note = "说话人未分离或未识别学生，学生专注度/师生匹配度仅给参考区间"

    course_label = meta.get("course_label") or meta.get("course_id", "未知课程")
    f = {
        "lesson": course_label,
        "date": (meta.get("start") or "")[:10],
        "time": _time_range(meta),
        "subject": meta.get("course_type") or course_label,
        "student": meta.get("student_name") or "未识别学生",
        "teacher": meta.get("teacher") or "老师",
        "location": meta.get("location") or "",
        "duration_min": duration_min,
        "n_turns": len(segs),
        "total_chars": total_chars,
        "chars_per_min": round(total_chars / duration_min, 1) if duration_min else 0,
        "separated": separated,
        "t_turns": len(t_segs),
        "s_turns": len(s_segs) if separated else 0,
        "s_chars": s_chars if separated else 0,
        "s_ratio": s_ratio,
        "avg_stu_len": (s_chars // len(s_segs)) if (separated and s_segs) else None,
        "alternations": alternations,
        "inter_density": inter_density,
        "n_question": n_question,
        "q_density": q_density,
        "answered": answered if separated else None,
        "response_rate": response_rate if separated else None,
        "n_long": len(longs),
        "longest_sec": longest_sec,
        "n_silence": len(sil),
        "silence_min": silence_min,
        "silence_ratio": silence_ratio,
        "seg_ratio": seg_ratio,
        "stu_seg_ratio": stu_seg_ratio,
        "balance": balance,
        "topic_count": topic_count,
        "teach_score": teach_score,
        "active_score": active_score,
        "focus_score": focus_score,
        "focus_est": focus_est,
        "match_score": match_score,
        "focus_note": focus_note,
        "insight": "",
    }
    f["insight"] = _insight(f)
    return f


def _time_range(meta: dict) -> str:
    try:
        s = meta.get("start", "")[11:16]
        e = meta.get("end", "")[11:16]
        return f"{s}-{e}"
    except Exception:  # noqa: BLE001
        return ""


def _thirds_ratio(segs: list[dict], duration_sec: float) -> list[float]:
    if not segs or not duration_sec:
        return [0.0, 0.0, 0.0]
    t0 = segs[0]["start_sec"] if segs else 0
    buckets = [0, 0, 0]
    for s in segs:
        rel = (s["start_sec"] - t0) / max(duration_sec, 1)
        buckets[min(int(rel * 3), 2)] += len(s["text"])
    total = sum(buckets) or 1
    return [round(b / total * 100, 1) for b in buckets]


def _balance(ratio: list[float] | None) -> int | None:
    if not ratio or not any(ratio):
        return None
    return max(0, round(100 - max(abs(r - 33.33) for r in ratio)))


def _teach_score(q_density, n_long, longest_sec, topic_count) -> int:
    q_part = min(q_density / 2.5, 1) * 100
    pace = max(0.0, 100 - n_long * 2 - max(0, longest_sec - 180) * 0.1)
    topic = min(topic_count / 60, 1) * 100
    return round(0.4 * q_part + 0.3 * pace + 0.3 * topic)


def _active_score(n_turns, duration_min, silence_ratio, balance) -> int:
    dens = min((n_turns / duration_min if duration_min else 0) / 3, 1) * 100
    silence_ctrl = max(0.0, 100 - silence_ratio * 2)
    return round(0.4 * dens + 0.3 * silence_ctrl + 0.3 * (balance or 0))


def _insight(f: dict) -> str:
    """规则化解读文案（与示例报告同风格，分号连接）。"""
    parts = []
    if f["separated"] and f["s_ratio"] is not None:
        r = f["s_ratio"]
        if 40 <= r <= 60:
            parts.append(f"师生发言占比均衡（学生 {r:.0f}%），参与度良好")
        elif r > 60:
            parts.append(f"学生发言占比偏高（{r:.0f}%），学生表达充分，注意老师引导不被淹没")
        else:
            parts.append(f"学生发言占比偏低（{r:.0f}%），建议增加提问与让学生复述的环节")
        rr = f["response_rate"]
        if rr is not None:
            if rr >= 90:
                parts.append(f"老师提问应答率 {rr}%，师生问答衔接流畅")
            elif rr >= 70:
                parts.append(f"老师提问应答率 {rr}%，问答衔接较好")
            else:
                parts.append(f"老师提问应答率 {rr}% 偏低，建议提问后多留等待时间并追问")
        m = f["match_score"]
        if m is not None:
            if m >= 85:
                parts.append(f"师生匹配度 {m}（高度契合）")
            elif m >= 70:
                parts.append(f"师生匹配度 {m}（契合良好）")
            else:
                parts.append(f"师生匹配度 {m}，建议关注师生节奏与互动方式")
    else:
        parts.append("说话人未分离，学生侧指标为参考区间")
    if f["teach_score"] >= 85:
        parts.append("老师授课分优秀，引导-节奏-覆盖均衡")
    elif f["teach_score"] >= 70:
        parts.append("老师授课分良好")
    else:
        parts.append("老师授课分偏低，建议提高问句密度、拆分长讲解段")
    if f["active_score"] >= 90:
        parts.append("课堂活跃度高，互动密集")
    elif f["active_score"] >= 75:
        parts.append("课堂活跃度良好")
    else:
        parts.append("课堂活跃度偏低，注意沉默间隙控制与节奏变化")
    if f["seg_ratio"] and f["seg_ratio"][2] < 25:
        parts.append("后段活跃度下滑明显，注意收尾阶段的节奏保持")
    return "；".join(parts)


# ---------------------------------------------------------------- 2. HTML 报告

_HTML_SHELL = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>课堂专注力报告 · 系统二</title>
<style>
__CSS__
</style>
</head>
<body>
<header><div class="hd-inner">
<h1>📊 课堂专注力报告</h1>
<div class="sub">教培提分双系统 · 系统二 ATTENTION（注意力系统）· 数据源：课堂录制转写（本地 DSP 麦克风 + 云端 AI pipeline）</div>
<span class="badge">__BADGE__</span>
</div></header>
<div class="wrap">
__SEPARATED_BAR__

<div class="kpis">
__KPIS__
</div>

<div class="card">
<h3><span class="dot"></span>本课评分对比（__LESSON__）</h3>
<div class="chart" id="chartCompare"></div>
</div>

<div class="card">
<h3><span class="dot"></span>课堂分段活跃度（前 / 中 / 后各 1/3）</h3>
<div class="chart" id="chartSeg"></div>
</div>

<div class="card">
<h3><span class="dot"></span>课次指标明细</h3>
<div class="tbl-wrap"><table>
<tr><th>课次</th><th>科目</th><th>时长</th><th>话轮</th><th>字数</th><th>问句</th><th>长讲解段</th><th>最长单段</th><th>沉默间隙</th><th>授课分</th><th>活跃分</th><th>前/中/后活跃%</th><th>学生专注</th></tr>
__DETAIL_ROW__
</table></div>
</div>

<div class="card">
<h3><span class="dot"></span>__LESSON__ <span class="hint">（__DATE__ __TIME__ · __DUR__ 分钟 · 学生 __STUDENT__）</span></h3>
<div class="kpis mini">
__MINI_KPIS__
</div>
<div class="grid2">
<div class="mini-chart" id="seg0"></div>
<div class="mini-chart" id="ring0"></div>
</div>
<div class="legendline">__LEGEND__</div>
<div class="note" style="margin-top:8px"><b>解读：</b>__INSIGHT__</div>
</div>

<div class="card">
<h3><span class="dot"></span>指标口径</h3>
<div class="note" style="margin:0"><b>评分公式（透明可解释）</b><br>
· 老师授课分 = 40%×问句引导（问句数/分钟，2.5/min 封顶）+ 30%×讲解节奏（长讲解段惩罚）+ 30%×主题覆盖<br>
· 课堂活跃分 = 40%×话轮密度（每分钟话轮，3/min 封顶）+ 30%×沉默控制（间隙&gt;20s 占比惩罚）+ 30%×分段均衡<br>
· 学生专注度 = 40%×参与 + 30%×应答率 + 30%×分段均衡 —— 需说话人分离数据，未分离时仅给参考区间<br>
· 师生匹配度 = 40%×互动同步（说话人交替密度）+ 30%×参与均衡 + 30%×内容契合<br>
· 沉默间隙 = 相邻话轮间隔 &gt; 20 秒的累计时长；长讲解段 = 单段 ≥90 秒 或 ≥220 字</div>
</div>

<footer>教培提分双系统 · 课堂专注力报告 V1.1 · 课堂录制系统自动生成</footer>
</div>

<script id="echarts-lib">__ECHARTS__</script>
<script>
var L=__DATA__;
__CHARTS__
</script>
</body>
</html>
"""


def _kpi(color, title, value, unit, desc) -> str:
    unit_html = f"<small> {unit}</small>" if unit else ""
    return (f'<div class="kpi {color}"><div class="t">{title}</div>'
            f'<div class="v">{value}{unit_html}</div><div class="d">{desc}</div></div>')


def _tag(score) -> str:
    if score is None:
        return '<span class="tag warn">—</span>'
    cls = "ok" if score >= 80 else ("warn" if score >= 60 else "bad")
    return f'<span class="tag {cls}">{score}</span>'


def render_focus_html(f: dict, out_path: Path) -> Path:
    """渲染单课专注力 HTML 报告（模板与 05-专注力报告.html 一致）。"""
    css = (ASSETS / "focus_style.css").read_text(encoding="utf-8")
    echarts = (ASSETS / "echarts.min.js").read_text(encoding="utf-8")
    charts = (ASSETS / "focus_charts.js").read_text(encoding="utf-8")

    badge = f"{f['date']} · {f['student']} · {f['subject']}（{f['student']}）"
    if f["separated"]:
        bar = (f'<div class="okbar"><b>✓ 说话人已分离：</b>本课转写成功识别'
               f'<b>老师（{f["teacher"]}）</b>与<b>学生（{f["student"]}）</b>双说话人，'
               f'学生专注度（{f["focus_score"]} 分）与师生匹配度（{f["match_score"]} 分）均为<b>真实计算</b>。</div>')
    else:
        bar = ('<div class="warnbar"><b>⚠ 说话人未完全分离：</b>未能识别学生身份，'
               '学生专注度与师生匹配度<b>仅给参考区间</b>；'
               '可先用「声纹注册工具」为该学生注册声纹，下次课程即可真实计算。</div>')

    kpis = "".join([
        _kpi("blue", "分析课次", 1, "节", f["date"]),
        _kpi("teal", "总时长", f["duration_min"], "min",
             f"约 {round((f['duration_min'] or 0) / 60, 1)} 小时"),
        _kpi("orange", "平均问句密度", f["q_density"], "/min", "老师引导强度"),
        _kpi("blue", "平均授课分", f["teach_score"], "", "引导+节奏+覆盖"),
        _kpi("teal", "平均活跃分", f["active_score"], "", "话轮+沉默+均衡"),
        _kpi("red", "师生匹配度", f["match_score"] if f["match_score"] is not None else "—",
             "", "同步+均衡+契合"),
        _kpi("orange", "学生专注度",
             f["focus_score"] if f["focus_score"] is not None else (f["focus_est"] or "—"),
             "", "参与+应答+均衡" if f["separated"] else "参考区间"),
        _kpi("blue", "学生发言占比",
             f["s_ratio"] if f["s_ratio"] is not None else "—", "%", "理想 40-60%"),
        _kpi("teal", "提问应答率",
             f["response_rate"] if f["response_rate"] is not None else "—", "%",
             f"{f['n_question']} 次提问"),
    ])

    seg = f["seg_ratio"] or [0, 0, 0]
    detail = (f"<tr><td><b>{f['lesson']}</b><br><span class=\"hint\">{f['date']} {f['time']}</span></td>"
              f"<td>{f['subject']}</td><td>{f['duration_min']}min</td>"
              f"<td>{f['n_turns']}</td><td>{f['total_chars']}</td>"
              f"<td>{f['n_question']}（{f['q_density']}/min）</td>"
              f"<td>{f['n_long']}</td><td>{f['longest_sec']}s</td>"
              f"<td>{f['silence_min']}min（{f['silence_ratio']}%）</td>"
              f"<td>{_tag(f['teach_score'])}</td><td>{_tag(f['active_score'])}</td>"
              f"<td>{seg[0]} / {seg[1]} / {seg[2]}</td>"
              f"<td>{f['focus_score'] if f['focus_score'] is not None else '—'}</td></tr>")

    mini = "".join([
        _kpi("blue", "老师授课分", f["teach_score"], "", "引导+节奏+覆盖"),
        _kpi("teal", "课堂活跃分", f["active_score"], "", "话轮+沉默+均衡"),
        _kpi("orange", "学生专注度",
             f["focus_score"] if f["focus_score"] is not None else (f["focus_est"] or "—"),
             "", "参与+应答+均衡" if f["separated"] else "参考区间"),
        _kpi("red", "师生匹配度", f["match_score"] if f["match_score"] is not None else "—",
             "", "同步+均衡+契合"),
        _kpi("blue", "学生发言占比", f["s_ratio"] if f["s_ratio"] is not None else "—",
             "%", "理想 40-60%"),
        _kpi("teal", "提问应答率", f["response_rate"] if f["response_rate"] is not None else "—",
             "%", f"{f['answered']}/{f['n_question']} 次" if f["separated"] else f"{f['n_question']} 次提问"),
    ])

    legend = (f"分段活跃度（前/中/后 {seg[0]}% / {seg[1]}% / {seg[2]}%）　·　"
              f"问句 {f['n_question']} 个（{f['q_density']}/min）　·　"
              f"长讲解 {f['n_long']} 段 · 最长 {f['longest_sec']}s　·　"
              f"沉默 {f['silence_min']}min（{f['silence_ratio']}%）")

    html = (_HTML_SHELL
            .replace("__CSS__", css)
            .replace("__ECHARTS__", echarts)
            .replace("__CHARTS__", charts)
            .replace("__DATA__", json.dumps([f], ensure_ascii=False))
            .replace("__BADGE__", badge)
            .replace("__SEPARATED_BAR__", bar)
            .replace("__KPIS__", kpis)
            .replace("__LESSON__", f["lesson"])
            .replace("__DETAIL_ROW__", detail)
            .replace("__MINI_KPIS__", mini)
            .replace("__LEGEND__", legend)
            .replace("__INSIGHT__", f["insight"])
            .replace("__DATE__", f["date"])
            .replace("__TIME__", f["time"])
            .replace("__DUR__", str(f["duration_min"]))
            .replace("__STUDENT__", f["student"]))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html, encoding="utf-8")
    log.info("专注力 HTML 已生成: %s", out_path)
    return out_path


# ---------------------------------------------------------------- 3. 专注力报告.xlsx

def write_focus_report_xlsx(f: dict, out_path: Path) -> Path:
    """生成 专注力报告.xlsx（4 个 sheet，结构同示例）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    navy = PatternFill("solid", fgColor="0F2B52")
    white_bold = Font(color="FFFFFF", bold=True, name="微软雅黑")
    title_font = Font(bold=True, size=13, color="0F2B52", name="微软雅黑")
    normal = Font(name="微软雅黑", size=10)
    wb = Workbook()

    # ---- sheet1 报告概览 ----
    ws = wb.active
    ws.title = "报告概览"
    ws["A1"] = "课堂专注力报告 · 报告概览"
    ws["A1"].font = title_font
    ws["A2"] = (f"数据源：课堂录制转写（{f['date']} · {f['student']} · {f['subject']}）"
                f"｜生成时间：{f['date']}")
    ws["A2"].font = normal
    ws.merge_cells("A4:F4")
    if f["separated"]:
        ws["A4"] = (f"说话人已分离：老师（{f['teacher']}）与学生（{f['student']}）双说话人识别成功，"
                    f"学生专注度（{f['focus_score']} 分）、师生匹配度（{f['match_score']} 分）、"
                    f"学生发言占比（{f['s_ratio']}%）、提问应答率（{f['response_rate']}%）均为真实计算。")
    else:
        ws["A4"] = ("说话人未完全分离：学生专注度、师生匹配度仅给参考区间；"
                    "为该学生注册声纹后下次课程可真实计算。")
    ws["A4"].font = normal
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    rows = [
        ("分析课次", 1, f["date"]),
        ("总时长(分钟)", f["duration_min"], f["lesson"]),
        ("平均问句密度(/min)", f["q_density"], "老师引导强度"),
        ("平均授课分", f["teach_score"], "引导+节奏+覆盖"),
        ("平均活跃分", f["active_score"], "话轮+沉默+均衡"),
        ("师生匹配度", f["match_score"] if f["match_score"] is not None else "—", "同步+均衡+契合"),
        ("学生专注度", f["focus_score"] if f["focus_score"] is not None else (f["focus_est"] or "—"),
         "参与+应答+均衡" if f["separated"] else "参考区间"),
    ]
    for c, h in enumerate(("指标", "数值", "说明"), start=1):
        cell = ws.cell(row=6, column=c, value=h)
        cell.fill = navy
        cell.font = white_bold
    for r, row in enumerate(rows, start=7):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v).font = normal
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 40

    # ---- sheet2 课次指标明细 ----
    ws2 = wb.create_sheet("课次指标明细")
    headers = ["课次", "日期", "时间", "科目", "时长(min)", "话轮数", "总字数", "每分钟字数",
               "问句数", "问句密度(/min)", "长讲解段", "最长单段(秒)", "沉默间隙数",
               "沉默时长(min)", "沉默占比(%)", "前段活跃(%)", "中段活跃(%)", "后段活跃(%)",
               "活跃均衡分", "主题覆盖数", "老师授课分", "课堂活跃分", "学生专注度",
               "参考区间", "学生发言占比(%)", "学生话轮数", "应答率(%)", "互动轮次", "师生匹配度"]
    for c, h in enumerate(headers, start=1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.fill = navy
        cell.font = white_bold
    seg = f["seg_ratio"] or [None, None, None]
    vals = [f["lesson"], f["date"], f["time"], f["subject"], f["duration_min"],
            f["n_turns"], f["total_chars"], f["chars_per_min"], f["n_question"],
            f["q_density"], f["n_long"], f["longest_sec"], f["n_silence"],
            f["silence_min"], f["silence_ratio"], seg[0], seg[1], seg[2],
            f["balance"], f["topic_count"], f["teach_score"], f["active_score"],
            f["focus_score"] if f["focus_score"] is not None else "数据受限",
            f["focus_est"] or "—", f["s_ratio"], f["s_turns"],
            f["response_rate"], f["alternations"],
            f["match_score"] if f["match_score"] is not None else "数据受限"]
    for c, v in enumerate(vals, start=1):
        ws2.cell(row=2, column=c, value=v).font = normal
    ws2.column_dimensions["A"].width = 30

    # ---- sheet3 分段活跃度 ----
    ws3 = wb.create_sheet("分段活跃度")
    for c, h in enumerate(("课次", "前 1/3 活跃(%)", "中 1/3 活跃(%)", "后 1/3 活跃(%)",
                           "均衡度(0-100)", "解读"), start=1):
        cell = ws3.cell(row=1, column=c, value=h)
        cell.fill = navy
        cell.font = white_bold
    if seg[0] is not None and seg[0] - seg[2] > 8:
        seg_note = "前段偏强，注意后段保持"
    elif seg[2] is not None and seg[2] - seg[0] > 8:
        seg_note = "后段偏强，前段可加快进入状态"
    else:
        seg_note = "三段活跃度均衡"
    for c, v in enumerate((f["lesson"], seg[0], seg[1], seg[2], f["balance"], seg_note), start=1):
        ws3.cell(row=2, column=c, value=v).font = normal
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["F"].width = 30

    # ---- sheet4 指标口径与建议 ----
    ws4 = wb.create_sheet("指标口径与建议")
    ws4["A1"] = "指标口径与使用说明"
    ws4["A1"].font = title_font
    notes = [
        ("一、数据管道（课堂录制系统 · 方案B）", 3),
        ("1. 本地 DSP 麦克风分段录制 → COS → ASR（说话人分离）→ 声纹 1:N → 结构化转写", 4),
        ("2. 专注力分析引擎解析转写：话轮、问句、长讲解、沉默间隙、分段活跃度、师生互动", 5),
        ("3. 输出：专注力报告.html（可视化）+ 本工作簿（Excel）+ 专注力数据录入表.xlsx（累计）", 6),
        ("二、评分公式（透明可解释）", 8),
        ("· 老师授课分 = 40%×问句引导(问句数/分钟, 2.5/min封顶) + 30%×讲解节奏(长讲解惩罚) + 30%×主题覆盖", 9),
        ("· 课堂活跃分 = 40%×话轮密度(每分钟话轮, 3/min封顶) + 30%×沉默控制(间隙>20s占比惩罚) + 30%×分段均衡", 10),
        ("· 学生专注度 = 40%×参与(学生发言占比接近50%最优) + 30%×应答率 + 30%×学生分段均衡", 11),
        ("· 师生匹配度 = 40%×互动同步(说话人交替密度) + 30%×参与均衡 + 30%×内容契合(应答成功率)", 12),
        ("· 沉默间隙 = 相邻话轮间隔>20秒的累计；长讲解段 = 单段≥90秒或≥220字", 13),
        ("三、数据质量（本课）", 15),
        (f"· {f['focus_note'] or ('本课转写已识别双说话人（老师：' + f['teacher'] + ' / 学生：' + f['student'] + '），学生专注度与师生匹配度为真实计算')}", 16),
        ("四、建议动作", 18),
        (f"· {f['insight']}", 19),
    ]
    for text, r in notes:
        ws4.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        cell = ws4.cell(row=r, column=1, value=text)
        cell.font = Font(name="微软雅黑", size=10, bold=text[0] in "一二三四")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", (28, 20, 20, 40)):
        ws4.column_dimensions[col].width = w

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log.info("专注力报告.xlsx 已生成: %s", out_path)
    return out_path


# ---------------------------------------------------------------- 4. 专注力数据录入表.xlsx（累计）

_ENTRY_HEADERS = ["序号", "课次", "日期", "上课时间", "学生姓名", "负责老师", "科目",
                  "上课地点", "课堂时长(分)", "话轮数", "总字数", "问句数", "问句密度(/min)",
                  "学生话轮数", "学生发言占比(%)", "应答率(%)", "互动轮次", "长讲解段",
                  "最长单段(秒)", "沉默时长(分)", "沉默占比(%)", "前段活跃(%)", "中段活跃(%)",
                  "后段活跃(%)", "主题覆盖数", "老师授课分", "课堂活跃分", "学生专注度",
                  "师生匹配度", "数据质量", "备注"]


def update_focus_entry_xlsx(f: dict, master_path: Path) -> Path:
    """向累计《专注力数据录入表.xlsx》追加一课；文件不存在则按模板新建。

    数据质量=受限 的行：学生专注度/匹配度写「数据受限」文本，
    AVERAGEIF 自动忽略，不影响平均分（与模板设计一致）。
    """
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    navy = PatternFill("solid", fgColor="0F2B52")
    white_bold = Font(color="FFFFFF", bold=True, name="微软雅黑")
    normal = Font(name="微软雅黑", size=10)

    if master_path.exists():
        wb = load_workbook(master_path)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "课堂数据录入"
        ws["A1"] = "课堂专注力数据录入 —— 一次课堂一条记录（唯一事实表，可追加多节课累计统计）"
        ws["A1"].font = Font(bold=True, size=12, color="0F2B52", name="微软雅黑")
        ws["A2"] = "由课堂录制系统自动追加；浅蓝列=自动计算可留空。"
        for c, h in enumerate(_ENTRY_HEADERS, start=1):
            cell = ws.cell(row=3, column=c, value=h)
            cell.fill = navy
            cell.font = white_bold
        dv = DataValidation(type="list", formula1='"已分离,受限,未录制"',
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add("AD4:AD500")
        _build_student_sheet(wb.create_sheet("学生专注力统计"))
        _build_teacher_sheet(wb.create_sheet("老师专注力统计"))
        _build_dict_sheet(wb.create_sheet("字段字典"))
        _build_help_sheet(wb.create_sheet("使用说明"))

    ws = wb["课堂数据录入"]
    row = ws.max_row + 1
    seg = f["seg_ratio"] or [None, None, None]
    limited = not f["separated"]
    vals = [
        row - 3, f["lesson"], f["date"], f["time"], f["student"], f["teacher"],
        f["subject"], f["location"], f["duration_min"], f["n_turns"], f["total_chars"],
        f["n_question"], f["q_density"], f["s_turns"], f["s_ratio"],
        f["response_rate"] if not limited else "数据受限",
        f["alternations"], f["n_long"], f["longest_sec"], f["silence_min"],
        f["silence_ratio"], seg[0], seg[1], seg[2], f["topic_count"],
        f["teach_score"], f["active_score"],
        f["focus_score"] if not limited else "数据受限",
        f["match_score"] if not limited else "数据受限",
        "受限" if limited else "已分离",
        f.get("focus_note") or "系统自动生成",
    ]
    for c, v in enumerate(vals, start=1):
        ws.cell(row=row, column=c, value=v).font = normal

    _ensure_stat_row(wb["学生专注力统计"], f["student"], f["subject"], kind="student")
    _ensure_stat_row(wb["老师专注力统计"], f["teacher"], None, kind="teacher")

    master_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(master_path)
    log.info("专注力数据录入表 已追加: %s（第 %d 行）", master_path, row)
    return master_path


def _build_student_sheet(ws) -> None:
    from openpyxl.styles import Font, PatternFill
    navy = PatternFill("solid", fgColor="0F2B52")
    ws["A1"] = "学生专注力统计 —— 一行一学生（自动汇总「课堂数据录入」）"
    ws["A1"].font = Font(bold=True, size=12, color="0F2B52", name="微软雅黑")
    ws["A2"] = "按学生姓名自动聚合；数据受限的行不参与平均。"
    for c, h in enumerate(("学生姓名", "科目", "课次数", "平均专注度", "平均活跃分",
                           "平均匹配度", "平均授课分", "平均应答率(%)", "平均发言占比(%)",
                           "最近专注度", "评价"), start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.fill = navy
        cell.font = Font(color="FFFFFF", bold=True, name="微软雅黑")


def _build_teacher_sheet(ws) -> None:
    from openpyxl.styles import Font, PatternFill
    navy = PatternFill("solid", fgColor="0F2B52")
    ws["A1"] = "老师专注力统计 —— 一行一老师（自动汇总，1对1 口径）"
    ws["A1"].font = Font(bold=True, size=12, color="0F2B52", name="微软雅黑")
    ws["A2"] = "按负责老师自动聚合：平均授课分反映课堂引导水平，平均匹配度反映师生契合度。"
    for c, h in enumerate(("负责老师", "课次数", "覆盖学生数", "平均授课分", "平均活跃分",
                           "平均匹配度", "平均应答率(%)", "平均专注度", "评价"), start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.fill = navy
        cell.font = Font(color="FFFFFF", bold=True, name="微软雅黑")


def _build_dict_sheet(ws) -> None:
    from openpyxl.styles import Font, PatternFill
    navy = PatternFill("solid", fgColor="0F2B52")
    ws["A1"] = "字段字典与口径说明（与分析引擎 focus_data.json 字段对齐）"
    ws["A1"].font = Font(bold=True, size=12, color="0F2B52", name="微软雅黑")
    for c, h in enumerate(("字段名", "类型", "来源", "口径 / 公式"), start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.fill = navy
        cell.font = Font(color="FFFFFF", bold=True, name="微软雅黑")
    rows = [
        ("课次", "文本", "飞书课表/引擎", "课堂标识，如「潘奕恺 · Alevel数学(M1)」"),
        ("日期 / 上课时间", "日期/文本", "飞书课表", "课堂日期与计划时段"),
        ("学生姓名 / 负责老师", "文本", "飞书课表/声纹", "1对1 统计单元"),
        ("课堂时长(分)", "数值", "转写时间戳", "末段结束 − 首段开始"),
        ("话轮数 / 总字数", "数值", "转写", "话轮=说话人连续发言段；字数=全部转写字符"),
        ("问句数 / 密度", "数值", "文本规则", "老师话轮中命中问句特征（？/吗/呢/怎么/多少…）"),
        ("学生话轮数", "数值", "说话人分离", "学生身份的发言段数（需分离数据）"),
        ("学生发言占比(%)", "数值", "说话人分离", "学生字数 ÷ 总字数；理想 40-60%"),
        ("应答率(%)", "数值", "说话人分离", "老师问句后 30 秒内下一段为学生回应的比例"),
        ("互动轮次", "数值", "说话人分离", "相邻话轮说话人交替次数"),
        ("长讲解段 / 最长单段", "数值", "老师话轮", "单段≥90秒或≥220字；最长单段秒数"),
        ("沉默时长/占比", "数值", "时间戳", "相邻话轮间隔>20秒的累计与占比"),
        ("前/中/后段活跃(%)", "数值", "时间三等分", "每 1/3 时段字数占比"),
        ("主题覆盖数", "数值", "文本规则", "编号/公式特征词（如 F743、9709/12）去重数"),
        ("老师授课分", "0-100", "引擎", "40%问句引导+30%讲解节奏+30%主题覆盖"),
        ("课堂活跃分", "0-100", "引擎", "40%话轮密度+30%沉默控制+30%分段均衡"),
        ("学生专注度", "0-100", "引擎", "40%参与+30%应答率+30%学生分段均衡（需分离）"),
        ("师生匹配度", "0-100", "引擎", "40%互动同步+30%参与均衡+30%内容契合（需分离）"),
        ("数据质量", "枚举", "引擎", "已分离 / 受限 / 未录制"),
    ]
    for r, row in enumerate(rows, start=3):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v).font = Font(name="微软雅黑", size=10)
    for col, w in zip("ABCD", (24, 10, 14, 60)):
        ws.column_dimensions[col].width = w


def _build_help_sheet(ws) -> None:
    from openpyxl.styles import Font
    ws["A1"] = "使用说明 —— 三步上手"
    ws["A1"].font = Font(bold=True, size=12, color="0F2B52", name="微软雅黑")
    lines = [
        "一、数据流（课堂录制系统自动生成）",
        "1. 每节课录制 → ASR+声纹 → 专注力引擎 focus_data.json → 本表「课堂数据录入」自动追加一行",
        "2. 说话人未分离的课：学生专注度/匹配度写「数据受限」，不影响老师授课分/活跃分统计",
        "二、持续积累与统计",
        "1. 「课堂数据录入」每节课一行，往下追加，历史数据不丢",
        "2. 「学生专注力统计」按学生自动聚合：课次数/平均专注度/匹配度/最近专注度",
        "3. 「老师专注力统计」按老师自动聚合：平均授课分/匹配度/覆盖学生数（1对1 口径）",
        "三、与系统一的衔接",
        "1. 本表 = 系统二（注意力系统）的「老师学生全量表」等价物",
        "2. 学生专注度/匹配度可回填到系统一学生画像（匹配度<60 建议关注）",
    ]
    r = 3
    for line in lines:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        cell = ws.cell(row=r, column=1, value=line)
        cell.font = Font(name="微软雅黑", size=10, bold=line[0] in "一二三")
        r += 1 if line[0] in "一二三" else 1


def _ensure_stat_row(ws, name: str, subject: str | None, kind: str) -> None:
    """统计表中若无此人则补一行公式（引用 课堂数据录入 固定区域）。"""
    if not name or name == "未识别学生":
        return
    for r in range(4, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == name:
            return
    r = ws.max_row + 1
    ws.cell(row=r, column=1, value=name)
    if kind == "student":
        ws.cell(row=r, column=2, value=subject or "")
        E = "'课堂数据录入'!$E$4:$E$500"
        ws.cell(row=r, column=3, value=f"=COUNTIF({E},$A{r})")
        for col, src in zip("DEFGHI", ("AB", "AA", "AC", "Z", "P", "O")):
            rng = f"'课堂数据录入'!${src}$4:${src}$500"
            ws.cell(row=r, column=ord(col) - 64,
                    value=f'=IFERROR(AVERAGEIF({E},$A{r},{rng}),"")')
        ws.cell(row=r, column=10,
                value=f'=IFERROR(LOOKUP(2,1/(({E}=$A{r})*({E}<>"")),'
                      f"'课堂数据录入'!$AB$4:$AB$500),\"\")")
        ws.cell(row=r, column=11,
                value=f'=IF(OR($D{r}="",$F{r}=""),"待积累",'
                      f'IF(AND($D{r}>=80,$F{r}>=80),"高专注·高契合",'
                      f'IF(AND($D{r}>=60,$F{r}>=60),"良好","建议关注")))')
    else:
        F = "'课堂数据录入'!$F$4:$F$500"
        E = "'课堂数据录入'!$E$4:$E$500"
        ws.cell(row=r, column=2, value=f"=COUNTIF({F},$A{r})")
        ws.cell(row=r, column=3,
                value=f'=IFERROR(SUMPRODUCT((({F}=$A{r})*({E}<>""))/'
                      f'COUNTIF({E},{E}&"")),0)')
        for col, src in zip("DEFGH", ("Z", "AA", "AC", "P", "AB")):
            rng = f"'课堂数据录入'!${src}$4:${src}$500"
            ws.cell(row=r, column=ord(col) - 64,
                    value=f'=IFERROR(AVERAGEIF({F},$A{r},{rng}),"")')
        ws.cell(row=r, column=9,
                value=f'=IF(OR($D{r}="",$F{r}=""),"待积累",'
                      f'IF(AND($D{r}>=75,$F{r}>=75),"优秀",'
                      f'IF(AND($D{r}>=60,$F{r}>=60),"良好","建议关注")))')
